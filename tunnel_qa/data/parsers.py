import csv
import html
import io
import json
import re
from pathlib import Path

from tunnel_qa.core.errors import AppError
from tunnel_qa.data.knowledge_chunking import prepare_knowledge_chunks
from tunnel_qa.data.pdf_layout import extract_page_text


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.splitlines())
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise AppError("UNSUPPORTED_ENCODING", "文件编码无法识别", field="file")


def parse_document(filename: str, content: bytes) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        return _normalize_text(_decode(content)), "text/plain"
    if suffix in {".html", ".htm"}:
        raw = _decode(content)
        without_scripts = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", raw, flags=re.I | re.S)
        text = re.sub(r"<[^>]+>", "\n", without_scripts)
        return _normalize_text(html.unescape(text)), "text/html"
    if suffix == ".json":
        data = json.loads(_decode(content))
        return _normalize_text(json.dumps(data, ensure_ascii=False, indent=2)), "application/json"
    if suffix == ".csv":
        reader = csv.reader(io.StringIO(_decode(content)))
        text = "\n".join(" | ".join(cell.strip() for cell in row) for row in reader)
        return _normalize_text(text), "text/csv"
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise AppError("PARSER_UNAVAILABLE", "XLSX解析组件未安装", field="file") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        try:
            for worksheet in workbook.worksheets:
                lines.append(f"[工作表: {worksheet.title}]")
                for row in worksheet.iter_rows(values_only=True):
                    lines.append(" | ".join("" if cell is None else str(cell).replace("\n", " ") for cell in row))
        finally:
            workbook.close()
        return _normalize_text("\n".join(lines)), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if suffix == ".pdf":
        # Reconstruct lines by page coordinates. Some PDFs store one glyph per
        # text operator; pypdf's default stream order then emits one character
        # per line, which is not a suitable paragraph boundary.
        try:
            import pymupdf
        except ImportError:
            pymupdf = None
        if pymupdf is not None:
            with pymupdf.open(stream=content, filetype="pdf") as document:
                if document.needs_pass:
                    raise AppError("ENCRYPTED_PDF", "PDF 已加密，请先解密再上传", field="file")
                text = "\n\n".join(
                    f"[第{i}页]\n{extract_page_text(page)}"
                    for i, page in enumerate(document, 1)
                )
            return _normalize_text(text), "application/pdf"
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise AppError("PARSER_UNAVAILABLE", "PDF解析组件未安装", field="file") from exc
        pages = PdfReader(io.BytesIO(content)).pages
        text = "\n\n".join(f"[第{i}页]\n{page.extract_text() or ''}" for i, page in enumerate(pages, 1))
        return _normalize_text(text), "application/pdf"
    if suffix == ".docx":
        try:
            from docx import Document
            from docx.oxml.ns import qn
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:
            raise AppError("PARSER_UNAVAILABLE", "DOCX解析组件未安装", field="file") from exc
        document = Document(io.BytesIO(content))
        lines = []
        # document.paragraphs omits tables. Walk body children to keep table rows
        # beside the surrounding explanatory paragraphs in their source order.
        for child in document.element.body:
            if child.tag == qn("w:p"):
                lines.append(Paragraph(child, document).text)
            elif child.tag == qn("w:tbl"):
                lines.append("")
                for row in Table(child, document).rows:
                    lines.append(" | ".join(cell.text.replace("\n", " ").strip() for cell in row.cells))
                lines.append("")
        text = "\n".join(lines)
        return _normalize_text(text), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    raise AppError(
        "UNSUPPORTED_FILE_TYPE",
        f"暂不支持文件类型: {suffix or '无扩展名'}",
        field="file",
    )


def split_text(text: str, chunk_size: int, overlap: int) -> list[tuple[str, str]]:
    """Compatibility wrapper retaining body sections, with safe boundaries."""
    prepared = prepare_knowledge_chunks("document.txt", text, chunk_size, overlap, mode="full")
    return [(chunk.heading, chunk.text) for chunk in prepared.chunks]


def split_parsed_document(
    filename: str, text: str, chunk_size: int, overlap: int
) -> list[tuple[str, str, int | None]]:
    """Legacy tuples expose the first page; the new API also exposes page_end."""
    prepared = prepare_knowledge_chunks(filename, text, chunk_size, overlap, mode="full")
    return [(chunk.heading, chunk.text, chunk.page_start) for chunk in prepared.chunks]
