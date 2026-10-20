import io
import json

import pytest
from docx import Document

from tunnel_qa.data.knowledge_chunking import prepare_knowledge_chunks
from tunnel_qa.data.parsers import parse_document, split_parsed_document, split_text


def prepare(text, *, filename="test.txt", size=180, overlap=30, mode="key"):
    return prepare_knowledge_chunks(filename, text, size, overlap, source_type="standard", mode=mode)


def body(result):
    return "\n".join(chunk.text for chunk in result.chunks)


def test_numbered_clauses_and_list_items_are_content_not_discarded_headings():
    text = (
        "# 施工要求\n3.1.1 开挖前必须检查围岩。\n"
        "3.1.2 作业人员不得进入危险区域。\n3.1.3 通风设备应保持运行。\n"
        "1、切断电源。\n2、撤离人员。\n220 V;\n接地电阻应符合规定。"
    )
    result = prepare(text)
    for line in text.replace("# ", "").splitlines():
        assert line in body(result)
    assert all(chunk.heading == "施工要求" for chunk in result.chunks)
    assert all("220 V" not in chunk.heading for chunk in result.chunks)


def test_short_definition_is_retained_even_without_space_after_clause_number():
    text = "2.1.9沙地sand land\n风力作用下沙粒堆积成近乎水平状的沙质地貌。"
    result = prepare(text)
    assert "2.1.9沙地sand land" in body(result)
    assert "风力作用下沙粒堆积" in body(result)
    assert result.report["excluded_count"] == 0


def test_pdf_repairs_chinese_wrapping_and_retains_scientific_units():
    result = prepare(
        "[第1页]\n３．１．１ 温 度 不 得 超 过 ３８℃，\n"
        "面积为１０m²时应保持通风。\n\ue000\n＿\n· 1 ·",
        filename="standard.pdf",
    )
    assert "3.1.1 温度不得超过 38℃，面积为10m²时应保持通风。" in body(result)
    assert "\ue000" not in body(result)
    assert "＿" not in body(result)
    assert "· 1 ·" not in body(result)
    assert result.report["excluded_by_reason"]["extraction_noise"] == 2
    assert result.chunks[0].page_start == result.chunks[0].page_end == 1


def test_only_edge_repeated_headers_are_removed_and_repeated_requirements_remain():
    text = "\n\n".join(
        f"[第{page}页]\n山东大学硕士学位论文\n3.1.1 作业前必须通风。\n"
        "正文介绍山东大学硕士学位论文的技术成果。\n"
        f"本页第{page}个独立试验结果。\n· {page} ·"
        for page in range(1, 4)
    )
    result = prepare(text, filename="thesis.pdf", size=500, overlap=0)
    assert not any(line == "山东大学硕士学位论文" for line in body(result).splitlines())
    assert body(result).count("3.1.1 作业前必须通风。") == 3
    assert body(result).count("正文介绍山东大学硕士学位论文的技术成果。") == 3
    assert result.report["excluded_by_reason"]["repeated_header_footer"] == 3
    assert result.report["excluded_by_reason"]["page_number"] == 3


def test_running_chapter_heading_is_kept_once_for_inherited_path():
    text = "\n".join(
        f"[第{page}页]\n第二章 安全措施\n第{page}项试验表明温度稳定。\n· {page} ·"
        for page in range(1, 4)
    )
    result = prepare(text, filename="thesis.pdf", size=500, overlap=0)
    assert body(result).count("第二章 安全措施") == 1
    assert all(chunk.heading == "第二章 安全措施" for chunk in result.chunks)


def test_key_selection_only_removes_explicit_sections_and_keeps_uncertain_body():
    text = (
        "# 目录\n第一章 绪论........1\n第二章 试验........4\n"
        "# 第一章 绪论\n参考文献可以用来复核研究的原始条件。\n"
        "## 1.1 试验方法\n试验温度为38℃，应记录其变化。\n"
        "# 参考文献\n[1] 某作者，某论文，2024。\n"
        "# 致谢\n感谢某单位提供帮助。\n"
        "# 附录\n传感器校准误差不得超过0.1℃。"
    )
    result = prepare(text)
    assert "参考文献可以用来复核研究的原始条件。" in body(result)
    assert "试验温度为38℃" in body(result)
    assert "传感器校准误差不得超过0.1℃。" in body(result)
    assert "感谢某单位" not in body(result)
    assert "某作者" not in body(result)
    assert "........1" not in body(result)
    assert set(result.report["excluded_by_reason"]) == {"contents", "references", "acknowledgements"}
    assert any("第一章 绪论 / 1.1 试验方法" == chunk.heading for chunk in result.chunks)
    assert "rule_based_selection_not_semantic_summarization" in result.report["warnings"]


@pytest.mark.parametrize("mode", ["full", "clean"])
def test_non_key_modes_keep_reference_and_acknowledgement_sections(mode):
    result = prepare("# 参考文献\n[1] 实验报告。\n# 致谢\n感谢提供仪器。", mode=mode)
    assert "[1] 实验报告。" in body(result)
    assert "感谢提供仪器。" in body(result)
    assert result.report["excluded_count"] == 0


def test_pdf_contents_does_not_swallow_body_after_damaged_chapter_title():
    result = prepare(
        "[第1页]\n目录\n1 绪论........1\n1.1 背景........2\n"
        "[第2页]\n1雜\n1.1背景与研究意义\n断层破碎带必须进行超前探测。\n"
        "[第3页]\n1.2现场试验\n试验中不得省略围岩监测。",
        filename="thesis.pdf",
    )
    assert "........" not in body(result)
    assert "断层破碎带必须进行超前探测。" in body(result)
    assert "试验中不得省略围岩监测。" in body(result)


@pytest.mark.parametrize("damaged_heading", ["料文献", "##文献"])
def test_ocr_reference_heading_requires_corroborating_bibliography(damaged_heading):
    refs = "\n".join(f"[{i}] 作者，试验报告[J]，2024，1：1-10。" for i in range(1, 5))
    result = prepare(
        "[第1页]\n结论\n支护必须与围岩条件匹配。\n" + damaged_heading + "\n" + refs
        + "\n[第2页]\n附录\n传感器误差不得超过0.1℃。", filename="thesis.pdf",
    )
    assert "支护必须与围岩条件匹配。" in body(result)
    assert "作者，试验报告" not in body(result)
    assert "传感器误差不得超过0.1℃。" in body(result)
    uncertain = prepare("[第1页]\n料文献\n材料检测强度不得低于20 MPa。", filename="thesis.pdf")
    assert "材料检测强度不得低于20 MPa。" in body(uncertain)


def test_publication_notices_are_filtered_without_matching_normal_keyword_mentions():
    result = prepare(
        "网络首发时间：2026年10月14日\n"
        "本实验分析出版声明对数据溯源的影响。\n"
        "# 版权声明\n版权所有，禁止复制。\n# 研究结果\n围岩强度降低10%。"
    )
    assert "网络首发时间" not in body(result)
    assert "本实验分析出版声明对数据溯源的影响。" in body(result)
    assert "围岩强度降低10%。" in body(result)
    assert result.report["excluded_by_reason"]["publication_notice"] == 3


def test_all_noise_is_explicitly_reported_and_samples_are_bounded():
    result = prepare("# 参考文献\n" + "\n".join(f"[{i}] 非正文参考资料。" for i in range(50)))
    assert not result.chunks
    assert result.report["excluded_count"] == 51
    assert len(result.report["excluded_samples"]) == 20
    assert result.report["needs_review"]
    assert "no_knowledge_content_after_selection" in result.report["warnings"]
    json.dumps(result.report, ensure_ascii=False)


def test_cross_page_unfinished_condition_and_value_stay_together_with_range():
    result = prepare(
        "[第1页]\n# 温度控制\n当作业面温度超过38℃时，应\n· 1 ·\n\n"
        "[第2页]\n持续通风并撤离非必要人员。\n· 2 ·",
        filename="test.pdf",
    )
    assert len(result.chunks) == 1
    chunk = result.chunks[0]
    assert "当作业面温度超过38℃时，应持续通风并撤离非必要人员。" in chunk.text
    assert (chunk.page_start, chunk.page_end) == (1, 2)
    assert "cross_page_continuation" in chunk.warnings


def test_complete_pages_and_different_chapters_are_not_merged():
    result = prepare(
        "[第1页]\n# 通风\n第一页的通风要求。\n"
        "[第2页]\n第二页的独立要求。\n# 温度\n温度不得超过38℃。",
        filename="test.pdf", size=500,
    )
    assert len(result.chunks) == 3
    assert [(chunk.page_start, chunk.page_end) for chunk in result.chunks] == [(1, 1), (2, 2), (2, 2)]
    assert [chunk.heading for chunk in result.chunks] == ["通风", "通风", "温度"]


def test_ambiguous_cross_page_boundary_is_not_assumed_to_be_continuation():
    result = prepare("[第1页]\n第一部分试验说明\n[第2页]\n另一个试验测得20℃。", filename="test.pdf")
    assert len(result.chunks) == 2
    assert [(chunk.page_start, chunk.page_end) for chunk in result.chunks] == [(1, 1), (2, 2)]


def test_sentence_overlap_never_begins_with_a_truncated_sentence():
    sentences = [f"第{i}项应检查围岩并记录温度。" for i in range(1, 10)]
    result = prepare("".join(sentences), size=50, overlap=20)
    assert len(result.chunks) > 1
    for chunk in result.chunks:
        assert len(chunk.text) <= 50
        assert all(line in sentences for line in chunk.text.splitlines())
    assert result.chunks[0].text.splitlines()[-1] == result.chunks[1].text.splitlines()[0]
    assert all(sentence in body(result) for sentence in sentences)


def test_long_single_sentence_is_not_silently_truncated_and_needs_review():
    text = "隧道温度监测数据" * 20
    result = prepare(text, size=50, overlap=20)
    assert "".join(chunk.text for chunk in result.chunks) == text
    assert all(len(chunk.text) <= 50 for chunk in result.chunks)
    assert all("hard_split_long_sentence" in chunk.warnings for chunk in result.chunks)
    assert result.report["needs_review"]


def test_table_rows_are_atomic_and_header_is_repeated():
    rows = [f"TEMP-{i} | {30+i}℃ | 10m² | 不得进入" for i in range(7)]
    result = prepare("测点 | 温度 | 面积 | 要求\n" + "\n".join(rows), size=95, overlap=20)
    assert len(result.chunks) > 1
    for chunk in result.chunks:
        assert chunk.kind == "table"
        assert chunk.text.startswith("测点 | 温度 | 面积 | 要求\n")
        assert len(chunk.text) <= 95
    for row in rows:
        assert body(result).splitlines().count(row) == 1


def test_oversize_table_row_preserves_numbers_units_and_emits_warning():
    row = "试验1 | 38℃ | " + "需要完整记录的条件" * 30 + " | 10m²"
    result = prepare("编号 | 温度 | 条件 | 面积\n" + row, size=100)
    assert row in body(result)
    assert "oversize_table_row" in result.report["warnings"]
    assert result.report["needs_review"]


def test_docx_parser_includes_tables_in_document_order():
    document = Document()
    document.add_paragraph("试验条件如下。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "温度", "面积"
    table.cell(1, 0).text, table.cell(1, 1).text = "38℃", "10m²"
    document.add_paragraph("应按上述条件实施监测。")
    stream = io.BytesIO()
    document.save(stream)
    text, _ = parse_document("data.docx", stream.getvalue())
    assert text.index("试验条件") < text.index("温度 | 面积") < text.index("38℃ | 10m²") < text.index("应按上述")
    assert "38℃ | 10m²" in body(prepare(text, filename="data.docx"))


def test_legacy_wrappers_preserve_clauses_and_default_to_full_mode():
    text = "3.1.1 应检测温度。\n3.1.2 不得关闭通风。\n# 参考文献\n[1] 试验资料。"
    chunks = split_text(text, 100, 20)
    joined = "\n".join(chunk[1] for chunk in chunks)
    assert "3.1.1 应检测温度。" in joined
    assert "3.1.2 不得关闭通风。" in joined
    assert "[1] 试验资料。" in joined
    pdf_chunks = split_parsed_document("test.pdf", "[第5页]\n" + text, 100, 20)
    assert all(chunk[2] == 5 for chunk in pdf_chunks)


def test_same_evidence_is_not_deduplicated_across_pages_or_documents():
    text = "[第1页]\n必须检测温度。\n[第2页]\n必须检测温度。"
    first = prepare(text, filename="one.pdf")
    second = prepare(text, filename="two.pdf")
    assert len(first.chunks) == len(second.chunks) == 2
    assert [chunk.page_start for chunk in first.chunks] == [1, 2]
    assert body(first).count("必须检测温度。") == 2


def test_empty_pdf_pages_require_ocr_review_instead_of_inventing_text():
    result = prepare("[第1页]\n\n[第2页]\n有效试验结果。\n[第3页]\n", filename="test.pdf")
    assert result.report["empty_text_pages"] == [1, 3]
    assert result.report["needs_review"]
    assert len(result.chunks) == 1 and result.chunks[0].page_start == 2


@pytest.mark.parametrize("size,overlap", [(0, 0), (-1, 0), (100, -1), (100, 100)])
def test_invalid_chunking_parameters_are_rejected(size, overlap):
    with pytest.raises(ValueError):
        prepare("正文。", size=size, overlap=overlap)


def test_invalid_selection_mode_is_rejected():
    with pytest.raises(ValueError, match="mode"):
        prepare("正文。", mode="summary")


def test_standalone_parameter_digits_are_not_removed_as_page_numbers():
    result = prepare("[第1页]\n温度不得超过\n38\n℃。\n面积为\n10\nm²。", filename="test.pdf")
    assert "38" in body(result) and "10" in body(result)
    assert "℃" in body(result) and "m²" in body(result)


def test_numeric_ocr_fragments_do_not_become_section_headings():
    result = prepare("[第1页]\n3.2.46 1试验方案设计\n实际测量条件为温度35℃，不得直接外推。", filename="test.pdf")
    assert all(chunk.heading == "正文" for chunk in result.chunks)
    assert "实际测量条件为温度35℃" in body(result)


def test_pdf_parser_reconstructs_positioned_text_and_preserves_page_markers():
    import pymupdf
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 100), "3.1.1 Keep ventilation active.")
    page.insert_text((72, 125), "3.1.2 Do not exceed 38 C.")
    raw = document.tobytes()
    document.close()
    text, _ = parse_document("positioned.pdf", raw)
    assert "[第1页]" in text
    result = prepare(text, filename="positioned.pdf")
    assert "3.1.1 Keep ventilation active." in body(result)
    assert "3.1.2 Do not exceed 38 C." in body(result)


def test_bibliography_page_requires_multiple_independent_signals():
    refs = "\n".join(f"[{i}] Author. Title [J]. Journal, 2024." for i in range(1, 5))
    result = prepare("[第1页]\n" + refs + "\n· 1 ·\n[第2页]\n温度不得超过38℃。", filename="test.pdf")
    assert "Journal" not in body(result)
    assert "温度不得超过38℃。" in body(result)
    assert result.report["excluded_by_reason"]["bibliography_page"] >= 4
    assert 1 not in result.report["low_text_pages"]


@pytest.mark.parametrize("reference_heading", ["参考文献\n", ""])
def test_dense_bibliography_does_not_remove_preceding_conclusion(reference_heading):
    refs = "\n".join(f"[{i}] Author. Title [J]. Journal, 2024." for i in range(1, 5))
    conclusion = "当围岩变形持续增大时，不得省略现场监测与复核。"
    result = prepare(
        "[第1页]\n1 试验方法\n试验记录用于现场复核。\n"
        "[第2页]\n2 结论\n" + conclusion + "\n" + reference_heading + refs,
        filename="test.pdf",
    )
    assert conclusion in body(result)
    assert "bibliography_page" not in result.report["excluded_by_reason"]
    if reference_heading:
        assert "Journal" not in body(result)


def test_bibliography_and_appendix_on_same_page_preserves_appendix_body():
    refs = "\n".join(f"[{i}] Author. Title [J]. Journal, 2024." for i in range(1, 5))
    result = prepare(
        "[第1页]\n参考文献\n" + refs + "\n附录\n复核时不得修改原始试验条件。",
        filename="test.pdf",
    )
    assert "复核时不得修改原始试验条件。" in body(result)
    assert "Journal" not in body(result)


def test_heading_only_key_document_is_not_an_independent_knowledge_fact():
    result = prepare("# 这里是一份只有标题没有任何有效正文的空文档")
    assert result.chunks == []
    assert result.report["needs_review"]


@pytest.mark.parametrize("sign", ["-", "+", "−", "﹣", "±", "∓", "–", "—"])
@pytest.mark.parametrize("filename", ["signs.pdf", "signs.txt"])
def test_standalone_numeric_sign_is_not_deleted_as_decoration(sign, filename):
    result = prepare(f"试验温度为\n{sign}\n5℃，不得用于其他工况。", filename=filename)
    # Do not require speculative dehyphenation or numeric reconstruction. The
    # original sign must survive regardless of the chosen line separator.
    assert sign in body(result)
    assert "5℃，不得用于其他工况。" in body(result)
    assert "extraction_noise" not in result.report["excluded_by_reason"]


@pytest.mark.parametrize("quantity", ["5 kJ", "5 kWh", "5 m²", "5 °C", "5 Ω", "5 mm/s", "5 kN/m²", "-5", "−5", "5e-3 kJ", "5 kg·m/s²"])
def test_scientific_quantity_is_body_not_a_standalone_heading(quantity):
    result = prepare(f"[第1页]\n试验记录如下。\n{quantity}", filename="units.pdf")
    assert quantity in body(result)
    assert "standalone_heading" not in result.report["excluded_by_reason"]
    assert all(chunk.heading == "正文" for chunk in result.chunks)


def test_successive_quantities_and_repeated_page_edge_measurements_are_kept():
    result = prepare("[第1页]\n试验能量为\n5 kJ\n6 kJ\n7 kJ", filename="units.pdf")
    assert all(value in body(result) for value in ["5 kJ", "6 kJ", "7 kJ"])
    repeated = prepare("\n".join(
        f"[第{page}页]\n5 kJ\n第{page}次试验仍需核对原始记录。"
        for page in range(1, 4)
    ), filename="units.pdf")
    assert body(repeated).count("5 kJ") == 3
    assert "repeated_header_footer" not in repeated.report["excluded_by_reason"]


def test_quantity_protection_does_not_disable_titles_clauses_or_visual_separator_cleanup():
    result = prepare(
        "[第1页]\n1 Introduction\n试验温度为-5℃，不得外推。\n"
        "---\n___\n———\n第二章 安全措施\n3.1.1 必须检查围岩。\n"
        "3.1.2 不得进入危险区域。",
        filename="titles.pdf",
    )
    assert [chunk.heading for chunk in result.chunks] == ["1 Introduction", "第二章 安全措施"]
    assert "-5℃" in body(result)
    assert "3.1.1 必须检查围岩。" in body(result)
    assert "3.1.2 不得进入危险区域。" in body(result)
    assert result.report["excluded_by_reason"]["extraction_noise"] == 3
